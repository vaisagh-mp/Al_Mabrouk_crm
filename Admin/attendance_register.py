from datetime import date, datetime, time, timedelta
from decimal import Decimal
from collections import defaultdict
import csv
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from django.utils import timezone
from django.db.models import Q
from django.core.exceptions import ObjectDoesNotExist
from django.core.paginator import Paginator
from django.contrib import messages
from django.urls import reverse
from django.http import HttpResponse

from Admin.models import Employee, Attendance, Leave, Holiday, DailyAttendanceStatus
from service_logs.models import EmployeeServiceLog

OFFICE_LOCATIONS = {'BDO', 'DMC', 'ABU'}
TRAVEL_LOCATIONS = {'OVERSEAS'}
# every other Attendance.location code counts as PORT

STATUS_OPTIONS = [
    'Office',
    'Vessel',
    'Office+Vessel',
    'Absent',
    'Annual Leave',
    'Sick Leave',
    'Personal Leave',
    'Training',
    'REST',
    'Sunday',
    'Holiday',
]

STATUS_CODE_TO_LABEL = {
    'OFFICE': 'Office',
    'VESSEL': 'Vessel',
    'OFFICE_VESSEL': 'Office+Vessel',
    'ABSENT': 'Absent',
    'ANNUAL_LEAVE': 'Annual Leave',
    'SICK_LEAVE': 'Sick Leave',
    'PERSONAL_LEAVE': 'Personal Leave',
    'TRAINING': 'Training',
    'REST': 'REST',
    'SUNDAY': 'Sunday',
    'HOLIDAY': 'Holiday',
}

STATUS_LABEL_TO_CODE = {v: k for k, v in STATUS_CODE_TO_LABEL.items()}


def _wo_number(project):
    """Safely return work_order_number for a project or None."""
    if not project:
        return None
    try:
        wo = getattr(project, 'workorder', None)
        return wo.work_order_number if wo else None
    except ObjectDoesNotExist:
        return None


def _to_local_dt(dt):
    if dt is None:
        return None
    if timezone.is_naive(dt):
        dt = timezone.make_aware(dt, timezone.get_current_timezone())
    return timezone.localtime(dt)


def _service_log_location(log):
    cat = log.project.category if log.project else None
    if cat == 'OVERSEAS':
        return 'Travel'
    if cat == 'AT_BERTH':
        return 'Port'
    if log.location_type == 'SAILING':
        return 'Sailing'
    if log.location_type == 'ANCHORAGE':
        return 'Anchorage'
    return 'Port'


def _attendance_location(att):
    if att.travel_in_time and not att.log_out_time:
        return 'Travel'
    loc = (att.location or '').strip().upper()
    if not loc:
        return 'Office'
    if loc in OFFICE_LOCATIONS:
        return 'Office'
    if loc in TRAVEL_LOCATIONS:
        return 'Travel'
    return 'Port'


def build_attendance_register(start_date, end_date, search="", status_filter=""):
    """
    Build daily attendance register for all active employees across [start_date, end_date].
    Uses a fixed number of queries (no N+1).
    All date logic uses local dates (timezone.localtime(att.login_time).date()).
    """
    if isinstance(start_date, str):
        start_date = datetime.strptime(start_date, '%Y-%m-%d').date()
    if isinstance(end_date, str):
        end_date = datetime.strptime(end_date, '%Y-%m-%d').date()

    if start_date > end_date:
        return []

    # 1. Fetch eligible employees
    employee_qs = Employee.objects.filter(
        user__is_active=True
    ).filter(
        Q(is_employee=True) | Q(is_manager=True) | Q(is_administration=True) | Q(is_hr=True)
    ).select_related('user')

    search_query = (search or "").strip()
    if search_query:
        employee_qs = employee_qs.filter(
            Q(user__username__icontains=search_query) |
            Q(user__first_name__icontains=search_query) |
            Q(user__last_name__icontains=search_query)
        )

    employees = list(employee_qs)
    if not employees:
        return []

    emp_ids = [e.id for e in employees]
    user_to_emp = {e.user_id: e for e in employees}
    tz = timezone.get_current_timezone()

    # Time boundaries for Attendance login_time
    start_dt = timezone.make_aware(datetime.combine(start_date, time.min), tz)
    end_dt = timezone.make_aware(datetime.combine(end_date, time.max), tz)

    # 2. Bulk load Attendances (exclude REJECTED)
    att_qs = Attendance.objects.filter(
        employee_id__in=emp_ids,
        login_time__gte=start_dt,
        login_time__lte=end_dt,
    ).exclude(
        status='REJECTED'
    ).select_related('employee__user', 'project__workorder', 'vessel')

    att_by_emp_date = defaultdict(list)
    for att in att_qs:
        if not att.login_time:
            continue
        local_date = timezone.localtime(att.login_time).date()
        if start_date <= local_date <= end_date:
            att_by_emp_date[(att.employee_id, local_date)].append(att)

    # 3. Bulk load Service logs
    slog_qs = EmployeeServiceLog.objects.filter(
        employee_id__in=emp_ids,
        date__gte=start_date,
        date__lte=end_date,
    ).select_related('project__workorder', 'time_log__work_order')

    slog_by_emp_date = defaultdict(list)
    for slog in slog_qs:
        slog_by_emp_date[(slog.employee_id, slog.date)].append(slog)

    # 4. Bulk load Approved Leaves
    leave_qs = Leave.objects.filter(
        user_id__in=[e.user_id for e in employees],
        approval_status='APPROVED',
        from_date__lte=end_date,
        to_date__gte=start_date,
    )

    leave_by_emp_date = {}
    for leave in leave_qs:
        emp = user_to_emp.get(leave.user_id)
        if not emp:
            continue
        curr_d = max(leave.from_date, start_date)
        last_d = min(leave.to_date, end_date)
        while curr_d <= last_d:
            leave_by_emp_date[(emp.id, curr_d)] = leave
            curr_d += timedelta(days=1)

    # 5. Bulk load Holidays
    holiday_qs = Holiday.objects.filter(
        date__gte=start_date,
        date__lte=end_date,
    )
    holidays_by_date = {h.date: h for h in holiday_qs}

    # 6. Bulk load DailyAttendanceStatus
    status_qs = DailyAttendanceStatus.objects.filter(
        employee_id__in=emp_ids,
        date__gte=start_date,
        date__lte=end_date,
    )
    daily_status_by_emp_date = {(s.employee_id, s.date): s for s in status_qs}

    today = timezone.localtime(timezone.now()).date()

    # Generate dates
    num_days = (end_date - start_date).days + 1
    date_list = [start_date + timedelta(days=i) for i in range(num_days)]

    rows = []

    for d in date_list:
        is_holiday_date = d in holidays_by_date
        is_sunday = (d.weekday() == 6)

        for emp in employees:
            override = daily_status_by_emp_date.get((emp.id, d))
            day_atts = att_by_emp_date.get((emp.id, d), [])
            day_logs = slog_by_emp_date.get((emp.id, d), [])
            leave = leave_by_emp_date.get((emp.id, d))

            has_pending = any(a.status == 'PENDING' for a in day_atts)
            attendance_ids = [a.id for a in day_atts]
            service_log_ids = [l.id for l in day_logs]
            override_id = override.id if override else None

            # Helper for vessel/logs data extraction
            logs_sorted = sorted(day_logs, key=lambda l: (l.start_time, l.id))
            atts_sorted = sorted(day_atts, key=lambda a: (a.login_time or datetime.min, a.id))

            has_work = bool(day_atts or day_logs)

            # Determine status based on priority rules
            rule_applied = None
            if override is not None:
                rule_applied = 1
            elif has_work:
                rule_applied = 2
            elif leave is not None:
                rule_applied = 3
            elif is_holiday_date:
                rule_applied = 4
            elif is_sunday:
                rule_applied = 5
            elif d > today:
                # Rule 6: The date is in the future -> skip the row
                continue
            else:
                rule_applied = 7

            is_off_day = False
            off_day_label = ""
            if has_work and (is_holiday_date or is_sunday):
                is_off_day = True
                if is_holiday_date:
                    off_day_label = "Holiday working"
                else:
                    off_day_label = "Sunday working"

            status_code = ""
            status_label = ""
            location_label = ""
            start = None
            end = None
            total_hours = Decimal('0.00')
            ot_hours = Decimal('0.00')
            vessel_job = ['-']
            service_ref = ['-']
            wo_number = ['-']

            if rule_applied == 1:
                status_code = override.status
                status_label = STATUS_CODE_TO_LABEL.get(override.status, override.status)
                location_label = dict(DailyAttendanceStatus.LOCATION_CHOICES).get(override.location, '') if override.location else ''
            elif rule_applied == 2:
                if day_logs and not day_atts:
                    status_code = 'VESSEL'
                    status_label = 'Vessel'
                elif day_atts and day_logs:
                    status_code = 'OFFICE_VESSEL'
                    status_label = 'Office+Vessel'
                else:
                    status_code = 'OFFICE'
                    status_label = 'Office'
            elif rule_applied == 3:
                leave_type = (leave.leave_type or '').strip().upper()
                if leave_type == 'ANNUAL LEAVE':
                    status_code = 'ANNUAL_LEAVE'
                    status_label = 'Annual Leave'
                elif leave_type == 'SICK LEAVE':
                    status_code = 'SICK_LEAVE'
                    status_label = 'Sick Leave'
                else:
                    status_code = 'PERSONAL_LEAVE'
                    status_label = 'Personal Leave'
                location_label = ''
            elif rule_applied == 4:
                status_code = 'HOLIDAY'
                status_label = 'Holiday'
                location_label = 'Holiday'
            elif rule_applied == 5:
                status_code = 'SUNDAY'
                status_label = 'Sunday'
                location_label = 'Holiday'
            elif rule_applied == 7:
                status_code = 'ABSENT'
                status_label = 'Absent'
                location_label = ''

            # Fill in work data if work exists (applies to rule 2 and rule 1 if work data exists)
            if has_work:
                if day_logs and not day_atts:
                    # Vessel work only
                    v_list = []
                    for l in logs_sorted:
                        v = l.vessel_name or (l.project.vessel_name if l.project else None) or (l.project.name if l.project else None)
                        if v and v.strip() and v.strip() not in v_list:
                            v_list.append(v.strip())
                    vessel_job = v_list if v_list else ['-']

                    r_list = []
                    for l in logs_sorted:
                        r = l.service_reference or (l.project.service_reference if l.project else None)
                        if r and r.strip() and r.strip() not in r_list:
                            r_list.append(r.strip())
                    service_ref = r_list if r_list else ['-']

                    w_list = []
                    for l in logs_sorted:
                        w = None
                        if l.time_log and l.time_log.work_order and l.time_log.work_order.work_order_number:
                            w = l.time_log.work_order.work_order_number
                        elif l.project:
                            w = _wo_number(l.project)
                        if w and w.strip() and w.strip() not in w_list:
                            w_list.append(w.strip())
                    wo_number = w_list if w_list else ['-']

                    longest_log = max(logs_sorted, key=lambda l: (l.total_hours or Decimal('0'), l.start_time))
                    if rule_applied == 2:
                        location_label = _service_log_location(longest_log)

                    start = min(l.start_time for l in logs_sorted)
                    if any(l.end_time is None for l in logs_sorted):
                        end = None
                    else:
                        end = max(l.end_time for l in logs_sorted)

                    total_hours = sum((l.total_hours or Decimal('0')) for l in logs_sorted).quantize(Decimal('0.01'))
                    ot_hours = sum((l.ot_hours or Decimal('0')) for l in logs_sorted).quantize(Decimal('0.01'))

                elif day_atts and not day_logs:
                    # Office work only
                    v_list = []
                    for a in atts_sorted:
                        v = (a.vessel.name if a.vessel else None) or (a.project.vessel_name if a.project else None) or (a.project.name if a.project else None)
                        if v and v.strip() and v.strip() not in v_list:
                            v_list.append(v.strip())
                    vessel_job = v_list if v_list else ['-']

                    r_list = []
                    for a in atts_sorted:
                        r = a.project.service_reference if a.project else None
                        if r and r.strip() and r.strip() not in r_list:
                            r_list.append(r.strip())
                    service_ref = r_list if r_list else ['-']

                    w_list = []
                    for a in atts_sorted:
                        w = _wo_number(a.project) if a.project else None
                        if w and w.strip() and w.strip() not in w_list:
                            w_list.append(w.strip())
                    wo_number = w_list if w_list else ['-']

                    if rule_applied == 2:
                        location_label = _attendance_location(atts_sorted[0])

                    valid_logins = [timezone.localtime(a.login_time) for a in atts_sorted if a.login_time]
                    start = min(t.time() for t in valid_logins) if valid_logins else None

                    if any(a.log_out_time is None for a in atts_sorted) or not any(a.log_out_time for a in atts_sorted):
                        end = None
                    else:
                        valid_logouts = [timezone.localtime(a.log_out_time) for a in atts_sorted if a.log_out_time]
                        end = max(t.time() for t in valid_logouts) if valid_logouts else None

                    total_hours = sum((a.total_hours_of_work or Decimal('0')) for a in atts_sorted).quantize(Decimal('0.01'))
                    ot_hours = max(total_hours - Decimal('8.00'), Decimal('0.00')).quantize(Decimal('0.01'))

                else:
                    # Office+Vessel work
                    v_list = []
                    for l in logs_sorted:
                        v = l.vessel_name or (l.project.vessel_name if l.project else None) or (l.project.name if l.project else None)
                        if v and v.strip() and v.strip() not in v_list:
                            v_list.append(v.strip())
                    vessel_job = v_list if v_list else ['-']

                    r_list = []
                    for l in logs_sorted:
                        r = l.service_reference or (l.project.service_reference if l.project else None)
                        if r and r.strip() and r.strip() not in r_list:
                            r_list.append(r.strip())
                    service_ref = r_list if r_list else ['-']

                    w_list = []
                    for l in logs_sorted:
                        w = None
                        if l.time_log and l.time_log.work_order and l.time_log.work_order.work_order_number:
                            w = l.time_log.work_order.work_order_number
                        elif l.project:
                            w = _wo_number(l.project)
                        if w and w.strip() and w.strip() not in w_list:
                            w_list.append(w.strip())
                    wo_number = w_list if w_list else ['-']

                    longest_log = max(logs_sorted, key=lambda l: (l.total_hours or Decimal('0'), l.start_time))
                    if rule_applied == 2:
                        location_label = _service_log_location(longest_log)

                    # Interval union
                    has_open = any(a.log_out_time is None for a in atts_sorted) or any(l.end_time is None for l in logs_sorted)
                    intervals = []
                    all_starts = []
                    all_ends = []

                    for a in atts_sorted:
                        if a.login_time:
                            login_dt = timezone.localtime(a.login_time)
                            all_starts.append(login_dt)
                            if a.log_out_time:
                                logout_dt = timezone.localtime(a.log_out_time)
                                all_ends.append(logout_dt)
                                if logout_dt > login_dt:
                                    intervals.append((login_dt, logout_dt))

                    for l in logs_sorted:
                        s_dt = timezone.make_aware(datetime.combine(d, l.start_time), tz)
                        all_starts.append(s_dt)
                        if l.end_time:
                            end_d = d + timedelta(days=1) if l.end_time < l.start_time else d
                            e_dt = timezone.make_aware(datetime.combine(end_d, l.end_time), tz)
                            all_ends.append(e_dt)
                            if e_dt > s_dt:
                                intervals.append((s_dt, e_dt))

                    if intervals:
                        sorted_ints = sorted(intervals, key=lambda x: x[0])
                        merged = []
                        for s_i, e_i in sorted_ints:
                            if not merged:
                                merged.append([s_i, e_i])
                            else:
                                if s_i <= merged[-1][1]:
                                    merged[-1][1] = max(merged[-1][1], e_i)
                                else:
                                    merged.append([s_i, e_i])
                        tot_secs = sum((m[1] - m[0]).total_seconds() for m in merged)
                        total_hours = Decimal(str(round(tot_secs / 3600.0, 2))).quantize(Decimal('0.01'))
                    else:
                        total_hours = Decimal('0.00')

                    ot_hours = max(total_hours - Decimal('8.00'), Decimal('0.00')).quantize(Decimal('0.01'))
                    start = min(all_starts).time() if all_starts else None
                    if has_open or not all_ends:
                        end = None
                    else:
                        end = max(all_ends).time()

            name = emp.user.get_full_name() or emp.user.username
            rank = emp.rank or "-"

            row = {
                'date': d,
                'employee': emp,
                'name': name,
                'rank': rank,
                'status_code': status_code,
                'status_label': status_label,
                'location_label': location_label,
                'vessel_job': vessel_job,
                'start': start,
                'end': end,
                'total_hours': total_hours,
                'ot_hours': ot_hours,
                'service_ref': service_ref,
                'wo_number': wo_number,
                'attendance_ids': attendance_ids,
                'service_log_ids': service_log_ids,
                'override_id': override_id,
                'has_pending': has_pending,
                'is_off_day': is_off_day,
                'off_day_label': off_day_label,
            }
            rows.append(row)

    if status_filter:
        flt = status_filter.strip().lower()
        rows = [r for r in rows if r['status_label'].lower() == flt or r['status_code'].lower() == flt]

    # Sort by date descending, then name ascending
    rows.sort(key=lambda r: (-r['date'].toordinal(), r['name'].lower()))
    return rows


STANDARD_DAY_HOURS = Decimal('8')


def build_attendance_summary(rows):
    """
    Groups attendance register rows by employee and produces summary metrics.
    Returns (summary_rows, totals).
    """
    emp_groups = defaultdict(list)
    emp_objects = {}
    emp_names = {}
    for r in rows:
        emp = r['employee']
        emp_id = emp.id
        emp_groups[emp_id].append(r)
        if emp_id not in emp_objects:
            emp_objects[emp_id] = emp
            emp_names[emp_id] = r['name']

    summary_rows = []
    for emp_id, emp in emp_objects.items():
        emp_rows = emp_groups[emp_id]
        name = emp_names[emp_id]
        designation = emp.rank if (emp.rank and emp.rank != '-') else emp.get_role()

        office_days = sum(1 for r in emp_rows if r['status_code'] == 'OFFICE')
        vessel_days = sum(1 for r in emp_rows if r['status_code'] == 'VESSEL')
        office_vessel_days = sum(1 for r in emp_rows if r['status_code'] == 'OFFICE_VESSEL')
        annual_leave_days = sum(1 for r in emp_rows if r['status_code'] == 'ANNUAL_LEAVE')
        sick_leave_days = sum(1 for r in emp_rows if r['status_code'] == 'SICK_LEAVE')
        personal_leave_days = sum(1 for r in emp_rows if r['status_code'] == 'PERSONAL_LEAVE')
        rest_days = sum(1 for r in emp_rows if r['status_code'] == 'REST')
        training_days = sum(1 for r in emp_rows if r['status_code'] == 'TRAINING')
        total_leave_days = annual_leave_days + sick_leave_days + personal_leave_days
        sunday_days = sum(1 for r in emp_rows if r['status_code'] == 'SUNDAY')
        holiday_days = sum(1 for r in emp_rows if r['status_code'] == 'HOLIDAY')

        total_hours = sum((Decimal(str(r.get('total_hours') or '0.00')) for r in emp_rows), Decimal('0.00')).quantize(Decimal('0.01'))
        gross_ot_hours = sum((Decimal(str(r.get('ot_hours') or '0.00')) for r in emp_rows), Decimal('0.00')).quantize(Decimal('0.01'))
        gross_ot_days = (gross_ot_hours / STANDARD_DAY_HOURS).quantize(Decimal('0.01'))
        absent_days = sum(1 for r in emp_rows if r['status_code'] == 'ABSENT')
        absent_hours = (Decimal(absent_days) * STANDARD_DAY_HOURS).quantize(Decimal('0.01'))
        off_day_worked = sum(1 for r in emp_rows if r.get('is_off_day') and r.get('status_code') not in ('SUNDAY', 'HOLIDAY'))

        summary_rows.append({
            'employee': emp,
            'name': name,
            'designation': designation,
            'office_days': office_days,
            'vessel_days': vessel_days,
            'office_vessel_days': office_vessel_days,
            'annual_leave_days': annual_leave_days,
            'sick_leave_days': sick_leave_days,
            'personal_leave_days': personal_leave_days,
            'rest_days': rest_days,
            'training_days': training_days,
            'total_leave_days': total_leave_days,
            'sunday_days': sunday_days,
            'holiday_days': holiday_days,
            'total_hours': total_hours,
            'gross_ot_hours': gross_ot_hours,
            'gross_ot_days': gross_ot_days,
            'absent_days': absent_days,
            'absent_hours': absent_hours,
            'off_day_worked': off_day_worked,
        })

    summary_rows.sort(key=lambda s: s['name'].lower())

    totals = {
        'name': 'TOTAL',
        'designation': '',
        'office_days': sum(r['office_days'] for r in summary_rows),
        'vessel_days': sum(r['vessel_days'] for r in summary_rows),
        'office_vessel_days': sum(r['office_vessel_days'] for r in summary_rows),
        'annual_leave_days': sum(r['annual_leave_days'] for r in summary_rows),
        'sick_leave_days': sum(r['sick_leave_days'] for r in summary_rows),
        'personal_leave_days': sum(r['personal_leave_days'] for r in summary_rows),
        'rest_days': sum(r['rest_days'] for r in summary_rows),
        'training_days': sum(r['training_days'] for r in summary_rows),
        'total_leave_days': sum(r['total_leave_days'] for r in summary_rows),
        'sunday_days': sum(r['sunday_days'] for r in summary_rows),
        'holiday_days': sum(r['holiday_days'] for r in summary_rows),
        'total_hours': sum((r['total_hours'] for r in summary_rows), Decimal('0.00')).quantize(Decimal('0.01')),
        'gross_ot_hours': sum((r['gross_ot_hours'] for r in summary_rows), Decimal('0.00')).quantize(Decimal('0.01')),
        'gross_ot_days': sum((r['gross_ot_days'] for r in summary_rows), Decimal('0.00')).quantize(Decimal('0.01')),
        'absent_days': sum(r['absent_days'] for r in summary_rows),
        'absent_hours': sum((r['absent_hours'] for r in summary_rows), Decimal('0.00')).quantize(Decimal('0.01')),
        'off_day_worked': sum(r['off_day_worked'] for r in summary_rows),
    }

    return summary_rows, totals


def summary_excel_response(summary_rows, totals, start, end):
    """
    Returns an HttpResponse with attendance summary Excel workbook (.xlsx).
    Filename: attendance_summary_<start>_<end>.xlsx
    """
    start_str = start.strftime('%Y-%m-%d') if hasattr(start, 'strftime') else str(start)
    end_str = end.strftime('%Y-%m-%d') if hasattr(end, 'strftime') else str(end)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance Summary"
    ws.views.sheetView[0].showGridLines = True

    # Theme Colors
    COLOR_PRIMARY = "1F3A52"
    COLOR_WORK = "2B5876"
    COLOR_LEAVE = "E67E22"
    COLOR_OFF = "6C757D"
    COLOR_HOURS = "3F6D95"
    COLOR_ABSENCE = "C0392B"
    COLOR_SUBHEADER = "E9EDF2"
    COLOR_TOTAL = "F1F5F9"
    COLOR_ZEBRA = "F8FAFC"
    COLOR_LEAVE_TINT = "FFF3CD"
    COLOR_ABSENT_TINT = "F8D7DA"

    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_subheader = Font(name="Calibri", size=10, bold=True, color="1A202C")
    font_data = Font(name="Calibri", size=10, color="1A202C")
    font_total = Font(name="Calibri", size=10, bold=True, color="1A202C")
    font_leave_tint = Font(name="Calibri", size=10, bold=True, color="856404")
    font_absent_tint = Font(name="Calibri", size=10, bold=True, color="721C24")

    fill_work = PatternFill(start_color=COLOR_WORK, end_color=COLOR_WORK, fill_type="solid")
    fill_leave = PatternFill(start_color=COLOR_LEAVE, end_color=COLOR_LEAVE, fill_type="solid")
    fill_off = PatternFill(start_color=COLOR_OFF, end_color=COLOR_OFF, fill_type="solid")
    fill_hours = PatternFill(start_color=COLOR_HOURS, end_color=COLOR_HOURS, fill_type="solid")
    fill_absence = PatternFill(start_color=COLOR_ABSENCE, end_color=COLOR_ABSENCE, fill_type="solid")
    fill_primary = PatternFill(start_color=COLOR_PRIMARY, end_color=COLOR_PRIMARY, fill_type="solid")
    fill_sub = PatternFill(start_color=COLOR_SUBHEADER, end_color=COLOR_SUBHEADER, fill_type="solid")
    fill_zebra = PatternFill(start_color=COLOR_ZEBRA, end_color=COLOR_ZEBRA, fill_type="solid")
    fill_total = PatternFill(start_color=COLOR_TOTAL, end_color=COLOR_TOTAL, fill_type="solid")
    fill_leave_tint = PatternFill(start_color=COLOR_LEAVE_TINT, end_color=COLOR_LEAVE_TINT, fill_type="solid")
    fill_absent_tint = PatternFill(start_color=COLOR_ABSENT_TINT, end_color=COLOR_ABSENT_TINT, fill_type="solid")

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    border_total = Border(
        left=thin_border_side, right=thin_border_side,
        top=Side(border_style="medium", color="3F6D95"),
        bottom=Side(border_style="double", color="3F6D95")
    )

    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 22

    # Row 1 headers
    ws.merge_cells("A1:A2")
    ws["A1"] = "Employee"
    ws["A1"].font = font_header
    ws["A1"].fill = fill_primary
    ws["A1"].alignment = align_center

    ws.merge_cells("B1:B2")
    ws["B1"] = "Designation"
    ws["B1"].font = font_header
    ws["B1"].fill = fill_primary
    ws["B1"].alignment = align_center

    groups = [
        ("C1:E1", "C1", "Work", fill_work),
        ("F1:K1", "F1", "Leave", fill_leave),
        ("L1:M1", "L1", "Off Days", fill_off),
        ("N1:P1", "N1", "Hours", fill_hours),
        ("Q1:R1", "Q1", "Absence", fill_absence),
    ]
    for cell_range, top_left, text, fill in groups:
        ws.merge_cells(cell_range)
        cell = ws[top_left]
        cell.value = text
        cell.font = font_header
        cell.fill = fill
        cell.alignment = align_center

    for col in range(1, 19):
        c1 = ws.cell(row=1, column=col)
        c1.border = border_cell
        if col in (1, 2):
            c1.fill = fill_primary
            ws.cell(row=2, column=col).fill = fill_primary
            ws.cell(row=2, column=col).border = border_cell
        elif 3 <= col <= 5:
            c1.fill = fill_work
        elif 6 <= col <= 11:
            c1.fill = fill_leave
        elif 12 <= col <= 13:
            c1.fill = fill_off
        elif 14 <= col <= 16:
            c1.fill = fill_hours
        elif 17 <= col <= 18:
            c1.fill = fill_absence

    sub_headers = [
        "Office Days", "Vessel Days", "Office+Vessel",
        "Annual Leave", "Sick Leave", "Personal Leave", "REST Days", "Training Days", "Total Leave Days",
        "Sunday", "Holiday",
        "Total Hours", "Gross OT Hrs", "Gross OT Days",
        "Absent Days", "Absent Hours"
    ]
    for idx, name in enumerate(sub_headers, start=3):
        cell = ws.cell(row=2, column=idx, value=name)
        cell.font = font_subheader
        cell.fill = fill_sub
        cell.alignment = align_center
        cell.border = border_cell

    current_row = 3
    for r_idx, r in enumerate(summary_rows):
        ws.row_dimensions[current_row].height = 20
        row_fill = fill_zebra if r_idx % 2 == 1 else None

        values = [
            (r['name'], align_left, font_data, None, None),
            (r['designation'], align_left, font_data, None, None),
            (r['office_days'], align_center, font_data, None, None),
            (r['vessel_days'], align_center, font_data, None, None),
            (r['office_vessel_days'], align_center, font_data, None, None),
            (r['annual_leave_days'], align_center, font_data, None, None),
            (r['sick_leave_days'], align_center, font_data, None, None),
            (r['personal_leave_days'], align_center, font_data, None, None),
            (r['rest_days'], align_center, font_data, None, None),
            (r['training_days'], align_center, font_data, None, None),
            (r['total_leave_days'], align_center, font_leave_tint if r['total_leave_days'] > 0 else font_data, fill_leave_tint if r['total_leave_days'] > 0 else None, None),
            (r['sunday_days'], align_center, font_data, None, None),
            (r['holiday_days'], align_center, font_data, None, None),
            (float(r['total_hours']), align_right, font_data, None, "0.00"),
            (float(r['gross_ot_hours']), align_right, font_data, None, "0.00"),
            (float(r['gross_ot_days']), align_right, font_data, None, "0.00"),
            (r['absent_days'], align_center, font_absent_tint if r['absent_days'] > 0 else font_data, fill_absent_tint if r['absent_days'] > 0 else None, None),
            (float(r['absent_hours']), align_right, font_absent_tint if r['absent_days'] > 0 else font_data, fill_absent_tint if r['absent_days'] > 0 else None, "0.00"),
        ]

        for col_idx, item in enumerate(values, start=1):
            val, alignment, font, custom_fill, num_format = item
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.alignment = alignment
            cell.font = font
            cell.border = border_cell
            if custom_fill:
                cell.fill = custom_fill
            elif row_fill:
                cell.fill = row_fill
            if num_format:
                cell.number_format = num_format

        current_row += 1

    # TOTAL row
    ws.row_dimensions[current_row].height = 22
    total_values = [
        ("TOTAL", align_left, font_total, fill_total, None),
        ("", align_left, font_total, fill_total, None),
        (totals['office_days'], align_center, font_total, fill_total, None),
        (totals['vessel_days'], align_center, font_total, fill_total, None),
        (totals['office_vessel_days'], align_center, font_total, fill_total, None),
        (totals['annual_leave_days'], align_center, font_total, fill_total, None),
        (totals['sick_leave_days'], align_center, font_total, fill_total, None),
        (totals['personal_leave_days'], align_center, font_total, fill_total, None),
        (totals['rest_days'], align_center, font_total, fill_total, None),
        (totals['training_days'], align_center, font_total, fill_total, None),
        (totals['total_leave_days'], align_center, font_total, fill_leave_tint if totals['total_leave_days'] > 0 else fill_total, None),
        (totals['sunday_days'], align_center, font_total, fill_total, None),
        (totals['holiday_days'], align_center, font_total, fill_total, None),
        (float(totals['total_hours']), align_right, font_total, fill_total, "0.00"),
        (float(totals['gross_ot_hours']), align_right, font_total, fill_total, "0.00"),
        (float(totals['gross_ot_days']), align_right, font_total, fill_total, "0.00"),
        (totals['absent_days'], align_center, font_total, fill_absent_tint if totals['absent_days'] > 0 else fill_total, None),
        (float(totals['absent_hours']), align_right, font_total, fill_absent_tint if totals['absent_days'] > 0 else fill_total, "0.00"),
    ]
    for col_idx, item in enumerate(total_values, start=1):
        val, alignment, font, custom_fill, num_format = item
        cell = ws.cell(row=current_row, column=col_idx, value=val)
        cell.alignment = alignment
        cell.font = font
        cell.fill = custom_fill
        cell.border = border_total
        if num_format:
            cell.number_format = num_format

    # Freeze panes at C3 (keeps Employee & Designation columns and Headers sticky)
    ws.freeze_panes = "C3"

    # Auto-fit column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="attendance_summary_{start_str}_{end_str}.xlsx"'
    return response


def summary_csv_response(summary_rows, totals, start, end):
    """
    Returns an HttpResponse with attendance summary CSV.
    Filename: attendance_summary_<start>_<end>.csv
    """
    start_str = start.strftime('%Y-%m-%d') if hasattr(start, 'strftime') else str(start)
    end_str = end.strftime('%Y-%m-%d') if hasattr(end, 'strftime') else str(end)
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="attendance_summary_{start_str}_{end_str}.csv"'

    writer = csv.writer(response)
    writer.writerow([
        "Employee",
        "Designation",
        "Office Days",
        "Vessel Days",
        "Office+Vessel",
        "Annual Leave",
        "Sick Leave",
        "Personal Leave",
        "REST Days",
        "Training Days",
        "Total Leave Days",
        "Sunday",
        "Holiday",
        "Total Hours",
        "Gross OT Hrs",
        "Gross OT Days",
        "Absent Days",
        "Absent Hours",
    ])

    for r in summary_rows:
        writer.writerow([
            r['name'],
            r['designation'],
            r['office_days'],
            r['vessel_days'],
            r['office_vessel_days'],
            r['annual_leave_days'],
            r['sick_leave_days'],
            r['personal_leave_days'],
            r['rest_days'],
            r['training_days'],
            r['total_leave_days'],
            r['sunday_days'],
            r['holiday_days'],
            str(r['total_hours']),
            str(r['gross_ot_hours']),
            str(r['gross_ot_days']),
            r['absent_days'],
            str(r['absent_hours']),
        ])

    writer.writerow([
        "TOTAL",
        "",
        totals['office_days'],
        totals['vessel_days'],
        totals['office_vessel_days'],
        totals['annual_leave_days'],
        totals['sick_leave_days'],
        totals['personal_leave_days'],
        totals['rest_days'],
        totals['training_days'],
        totals['total_leave_days'],
        totals['sunday_days'],
        totals['holiday_days'],
        str(totals['total_hours']),
        str(totals['gross_ot_hours']),
        str(totals['gross_ot_days']),
        totals['absent_days'],
        str(totals['absent_hours']),
    ])

    return response


def _register_context(request):
    """
    Shared context builder for admin and HR attendance register views.
    Handles start/end clamping, search, status filtering, summary counts,
    and pagination (10 per page).
    """
    today = timezone.localtime(timezone.now()).date()
    default_start = today.replace(day=1)
    default_end = today

    start_param = request.GET.get('start', '').strip()
    end_param = request.GET.get('end', '').strip()

    try:
        start_date = datetime.strptime(start_param, '%Y-%m-%d').date() if start_param else default_start
    except (ValueError, TypeError):
        start_date = default_start

    try:
        end_date = datetime.strptime(end_param, '%Y-%m-%d').date() if end_param else default_end
    except (ValueError, TypeError):
        end_date = default_end

    if start_date > end_date:
        start_date, end_date = end_date, start_date

    if (end_date - start_date).days > 62:
        messages.warning(request, "Date range cannot exceed 62 days. Range has been adjusted.")
        end_date = start_date + timedelta(days=62)

    search_query = request.GET.get('search', '').strip()
    status_filter = request.GET.get('status', '').strip()

    # Build all rows for this date and search range to compute accurate summary counts
    all_rows = build_attendance_register(start_date, end_date, search=search_query, status_filter="")

    summary_counts = {
        'office': sum(1 for r in all_rows if r['status_label'] == 'Office'),
        'vessel': sum(1 for r in all_rows if r['status_label'] == 'Vessel'),
        'office_vessel': sum(1 for r in all_rows if r['status_label'] == 'Office+Vessel'),
        'absent': sum(1 for r in all_rows if r['status_label'] == 'Absent'),
        'leave': sum(1 for r in all_rows if r['status_label'] in ('Annual Leave', 'Sick Leave', 'Personal Leave')),
        'training': sum(1 for r in all_rows if r['status_label'] == 'Training'),
        'rest': sum(1 for r in all_rows if r['status_label'] == 'REST'),
    }

    summary_rows_all, summary_totals = build_attendance_summary(all_rows)

    if len(summary_rows_all) > 50:
        sum_paginator = Paginator(summary_rows_all, 50)
        sum_page_num = request.GET.get('sum_page')
        summary_rows = sum_paginator.get_page(sum_page_num)
        summary_page_obj = summary_rows
    else:
        summary_rows = summary_rows_all
        summary_page_obj = None

    if status_filter:
        flt = status_filter.strip().lower()
        display_rows = [r for r in all_rows if r['status_label'].lower() == flt or r['status_code'].lower() == flt]
    else:
        display_rows = all_rows

    paginator = Paginator(display_rows, 10)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    query_params = request.GET.copy()
    if 'page' in query_params:
        del query_params['page']
    querystring_without_page = query_params.urlencode()

    return {
        'rows': page_obj,
        'page_obj': page_obj,
        'start': start_date.strftime('%Y-%m-%d'),
        'end': end_date.strftime('%Y-%m-%d'),
        'search_query': search_query,
        'status_filter': status_filter,
        'status_options': STATUS_OPTIONS,
        'summary_counts': summary_counts,
        'summary_rows': summary_rows,
        'summary_rows_all': summary_rows_all,
        'summary_totals': summary_totals,
        'summary_page_obj': summary_page_obj,
        'summary_range_label': f"{start_date.strftime('%d %b %Y')} – {end_date.strftime('%d %b %Y')}",
        'querystring_without_page': querystring_without_page,
    }


def handle_set_day_status(request, employee_id, date_str, role='admin'):
    """
    Shared view handler to create, update, or clear DailyAttendanceStatus.
    """
    from django.shortcuts import get_object_or_404, redirect, render
    from django.http import Http404

    employee = get_object_or_404(Employee, id=employee_id)
    try:
        status_date = datetime.strptime(date_str, '%Y-%m-%d').date()
    except (ValueError, TypeError):
        raise Http404("Invalid date format")

    override = DailyAttendanceStatus.objects.filter(employee=employee, date=status_date).first()

    fallback_url_name = 'attendance_list_adminview' if role == 'admin' else 'hr_attendance_list_view'
    next_url = request.GET.get('next') or request.POST.get('next')
    if not next_url or not next_url.startswith('/') or next_url.startswith('//'):
        next_url = reverse(fallback_url_name)

    if request.method == 'POST':
        action = request.POST.get('action')
        if action == 'clear' or 'clear' in request.POST:
            if override:
                override.delete()
                messages.success(request, f"Manual status for {employee.user.get_full_name() or employee.user.username} on {status_date} cleared.")
            else:
                messages.info(request, "No manual status to clear.")
            return redirect(next_url)

        status_val = request.POST.get('status')
        valid_statuses = [code for code, label in DailyAttendanceStatus.STATUS_CHOICES]
        if status_val not in valid_statuses:
            messages.error(request, "Invalid status choice selected.")
        else:
            loc_val = request.POST.get('location') or None
            valid_locations = [code for code, label in DailyAttendanceStatus.LOCATION_CHOICES]
            if loc_val and loc_val not in valid_locations:
                loc_val = None
            note_val = request.POST.get('note', '').strip()

            if override:
                override.status = status_val
                override.location = loc_val
                override.note = note_val
                override.updated_by = request.user
                override.save()
            else:
                DailyAttendanceStatus.objects.create(
                    employee=employee,
                    date=status_date,
                    status=status_val,
                    location=loc_val,
                    note=note_val,
                    updated_by=request.user
                )
            messages.success(request, f"Manual status for {employee.user.get_full_name() or employee.user.username} on {status_date} saved successfully.")
            return redirect(next_url)

    base_template = 'Admin/adminbase.html' if role == 'admin' else 'hr/hrbase.html'
    return render(request, 'attendance/set_day_status.html', {
        'employee': employee,
        'status_date': status_date,
        'date_str': date_str,
        'override': override,
        'status_choices': DailyAttendanceStatus.STATUS_CHOICES,
        'location_choices': DailyAttendanceStatus.LOCATION_CHOICES,
        'next_url': next_url,
        'role': role,
        'base_template': base_template,
    })
